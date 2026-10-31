from app.tools.docx_structure import split_document_structure


def test_short_document_keeps_every_source_block_without_h2():
    source = "# Report\n\nFirst point.\nSecond point.\n\n| Metric | Value |\n| --- | --- |\n| Count | 73421 |"
    title, abstract, body, has_toc = split_document_structure(
        source.splitlines(), cover_title="Report",
        toc_headings=["Contents"],
    )
    assert title == "Report"
    assert abstract == ""
    assert has_toc is False
    assert "\n".join(body) == source


def test_word_export_contains_short_document_body_and_table():
    from pathlib import Path

    from docx import Document

    from app.tools.docx import generate_docx_file

    source = (
        "# Release report\n\nFirst point.\nSecond point.\n\n"
        "| Metric | Value |\n| --- | --- |\n| Count | 73421 |"
    )
    path = Path(generate_docx_file(source, "test_docx_structure_export.docx", cover_title="Release report"))
    try:
        document = Document(path)
        contents = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        assert "First point." in contents
        assert "Second point." in contents
        assert "73421" in contents
    finally:
        path.unlink(missing_ok=True)
